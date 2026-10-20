"""Save contacts to an Excel sheet, updating an existing row when the email or phone matches."""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADER = ["Name", "Company", "Title", "Email", "Phone", "Created_at", "Updated_at", "Source"]


def _digits(s: str | None) -> str:
    return re.sub(r"\D", "", s or "")


def upsert_contact(path: str | Path, fields: dict, source: str) -> str:
    path = Path(path)
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.append(HEADER)
    ws = wb.active
    email, phone = (fields.get("Email") or "").lower(), _digits(fields.get("Phone"))
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for row in ws.iter_rows(min_row=2):
        same_email = email and (row[3].value or "").lower() == email
        same_phone = phone and _digits(row[4].value) == phone
        if same_email or same_phone:
            for col, key in enumerate(HEADER[:5]):
                if not row[col].value and fields.get(key):
                    row[col].value = fields[key]          # only fill gaps, never overwrite
            row[6].value = now
            wb.save(path)
            return "updated"
    ws.append([fields.get(k, "") for k in HEADER[:5]] + [now, now, source])
    wb.save(path)
    return "inserted"
